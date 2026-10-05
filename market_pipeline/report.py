"""Daily market report: breadth, top movers, most active, 5-day momentum -> Excel + chart + HTML email."""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import pandas as pd  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from .config import MIN_TURNOVER_FOR_RANKING, RANKING_SERIES, REPORT_DIR  # noqa: E402
from .storage import query  # noqa: E402

UP, DOWN, NEUTRAL = "#2a78d6", "#e34948", "#b8b6ae"     # diverging poles + neutral
INK, INK_2, SURFACE, GRID = "#0b0b0b", "#52514e", "#fcfcfb", "#e6e4df"
HEADER = PatternFill("solid", fgColor="1F3864")


@dataclass
class DailyReport:
    trade_date: str
    breadth: dict
    gainers: pd.DataFrame
    losers: pd.DataFrame
    active: pd.DataFrame
    momentum: pd.DataFrame
    quality: dict

    def headline(self) -> str:
        b = self.breadth
        text = (f"{b['advances']} advances vs {b['declines']} declines (A/D ratio {b['ad_ratio']}), "
                f"total turnover INR {b['turnover_cr']:,.0f} cr.")
        if len(self.gainers) and len(self.losers):
            g, l = self.gainers.iloc[0], self.losers.iloc[0]
            text += f" Top gainer {g.symbol} {g.change_pct:+.2f}%, top loser {l.symbol} {l.change_pct:+.2f}%."
        return text


def build(trade_date: str, top_n: int = 10) -> DailyReport:
    series = ",".join(f"'{s}'" for s in RANKING_SERIES)
    day = query(f"SELECT * FROM daily_prices WHERE trade_date = ? AND series IN ({series})", (trade_date,))
    if day.empty:
        raise ValueError(f"No data loaded for {trade_date}")
    liquid = day[day.turnover >= MIN_TURNOVER_FOR_RANKING]
    cols = ["symbol", "close", "change_pct", "volume", "turnover"]
    adv, dec = int((day.change_pct > 0).sum()), int((day.change_pct < 0).sum())
    breadth = {"advances": adv, "declines": dec, "unchanged": int((day.change_pct == 0).sum()),
               "ad_ratio": round(adv / dec, 2) if dec else None,
               "turnover_cr": round(day.turnover.sum() / 1e7, 1), "symbols": int(len(day))}

    dates = query("SELECT DISTINCT trade_date FROM daily_prices WHERE trade_date <= ? "
                  "ORDER BY trade_date DESC LIMIT 6", (trade_date,)).trade_date.tolist()
    momentum = pd.DataFrame(columns=["symbol", "close_5d_ago", "close", "change_5d_pct"])
    if len(dates) == 6:
        hist = query(f"SELECT symbol, trade_date, close FROM daily_prices WHERE trade_date IN (?, ?) "
                     f"AND series IN ({series})", (dates[-1], trade_date))
        piv = hist.pivot_table(index="symbol", columns="trade_date", values="close").dropna()
        piv.columns = ["close_5d_ago", "close"]
        piv["change_5d_pct"] = ((piv.close / piv.close_5d_ago - 1) * 100).round(2)
        momentum = piv.reset_index().sort_values("change_5d_pct", ascending=False).head(top_n)

    q = query("SELECT rows_in, rows_loaded, rows_rejected, issues FROM pipeline_runs "
              "WHERE trade_date = ? AND status = 'ok' ORDER BY run_at DESC LIMIT 1", (trade_date,))
    return DailyReport(
        trade_date=trade_date, breadth=breadth,
        gainers=liquid.nlargest(top_n, "change_pct")[cols].reset_index(drop=True),
        losers=liquid.nsmallest(top_n, "change_pct")[cols].reset_index(drop=True),
        active=day.nlargest(top_n, "turnover")[cols].reset_index(drop=True),
        momentum=momentum.reset_index(drop=True),
        quality=q.iloc[0].to_dict() if len(q) else {},
    )


def chart(rep: DailyReport, path: Path) -> Path:
    movers = pd.concat([rep.gainers.head(8), rep.losers.head(8).iloc[::-1]]).sort_values("change_pct")
    fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(14, 7), gridspec_kw={"width_ratios": [2.2, 1]},
                                   facecolor=SURFACE)
    for ax in (ax1, ax2):
        ax.set_facecolor(SURFACE)
        for s in ("top", "right", "left"):
            ax.spines[s].set_visible(False)
        ax.spines["bottom"].set_color(GRID)
        ax.tick_params(colors=INK_2, labelsize=11, length=0)

    colors = [UP if v > 0 else DOWN for v in movers.change_pct]
    ax1.barh(movers.symbol, movers.change_pct, color=colors, height=0.62, edgecolor=SURFACE, linewidth=2)
    ax1.axvline(0, color=INK_2, linewidth=1)
    ax1.xaxis.grid(True, color=GRID, linewidth=0.8)
    ax1.set_axisbelow(True)
    for y, v in enumerate(movers.change_pct):
        ax1.text(v + (0.15 if v > 0 else -0.15), y, f"{v:+.2f}%", va="center",
                 ha="left" if v > 0 else "right", fontsize=10, color=INK)
    lim = max(abs(movers.change_pct).max() * 1.3, 1)
    ax1.set_xlim(-lim, lim)
    ax1.set_title("Top gainers and losers (% change, liquid stocks)", loc="left", fontsize=14, color=INK, pad=12)

    b = rep.breadth
    vals = [b["advances"], b["unchanged"], b["declines"]]
    ax2.bar(["Advances", "Unchanged", "Declines"], vals, color=[UP, NEUTRAL, DOWN], width=0.6,
            edgecolor=SURFACE, linewidth=2)
    for x, v in enumerate(vals):
        ax2.text(x, v + max(vals) * 0.02, str(v), ha="center", fontsize=12, color=INK)
    ax2.set_yticks([])
    ax2.set_title(f"Market breadth (A/D {b['ad_ratio']})", loc="left", fontsize=14, color=INK, pad=12)

    fig.suptitle(f"Market summary - {rep.trade_date}", x=0.01, ha="left", fontsize=18, color=INK, weight="bold")
    fig.tight_layout(rect=(0, 0, 1, 0.95))
    fig.savefig(path, dpi=110, facecolor=SURFACE)
    plt.close(fig)
    return path


def excel(rep: DailyReport, path: Path) -> Path:
    summary = pd.DataFrame([
        ("Trade date", rep.trade_date), ("Symbols", rep.breadth["symbols"]),
        ("Advances", rep.breadth["advances"]), ("Declines", rep.breadth["declines"]),
        ("Unchanged", rep.breadth["unchanged"]), ("A/D ratio", rep.breadth["ad_ratio"]),
        ("Turnover (INR cr)", rep.breadth["turnover_cr"]),
        ("Rows received", rep.quality.get("rows_in")), ("Rows loaded", rep.quality.get("rows_loaded")),
        ("Rows quarantined", rep.quality.get("rows_rejected")), ("Quality issues", rep.quality.get("issues")),
    ], columns=["Metric", "Value"])
    sheets = {"Summary": summary, "Top Gainers": rep.gainers, "Top Losers": rep.losers,
              "Most Active": rep.active, "Momentum 5D": rep.momentum}
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
    wb = load_workbook(path)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        for c in ws[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), HEADER
            c.alignment = Alignment(horizontal="center")
        for i in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(i)].width = 22
        headers = [c.value for c in ws[1]]
        for row in ws.iter_rows(min_row=2):
            for c in row:
                h = headers[c.column - 1]
                if h in ("turnover",):
                    c.number_format = "#,##0"
                elif h in ("change_pct", "change_5d_pct"):
                    c.number_format = '+0.00"%";-0.00"%"'
                    c.font = Font(color="2A78D6" if (c.value or 0) > 0 else "E34948")
    wb.save(path)
    return path


def email_html(rep: DailyReport) -> str:
    def table(df, title):
        rows = "".join(
            f"<tr><td style='padding:4px 10px'>{r.symbol}</td><td style='padding:4px 10px;text-align:right'>"
            f"{r.close:,.2f}</td><td style='padding:4px 10px;text-align:right;color:{UP if r.change_pct > 0 else DOWN}'>"
            f"{r.change_pct:+.2f}%</td></tr>" for r in df.head(5).itertuples())
        return (f"<h3 style='margin:18px 0 6px'>{title}</h3><table style='border-collapse:collapse'>"
                f"<tr style='background:#1F3864;color:#fff'><th style='padding:4px 10px'>Symbol</th>"
                f"<th style='padding:4px 10px'>Close</th><th style='padding:4px 10px'>Change</th></tr>{rows}</table>")
    return (f"<div style='font-family:Segoe UI,Arial,sans-serif;font-size:14px;color:#222'>"
            f"<p><b>Market summary for {rep.trade_date}</b></p><p>{rep.headline()}</p>"
            f"{table(rep.gainers, 'Top gainers')}{table(rep.losers, 'Top losers')}"
            f"<p style='color:#666;font-size:12px;margin-top:18px'>Full report attached. "
            f"Data quality: {rep.quality.get('rows_loaded', '?')} rows loaded, "
            f"{rep.quality.get('rows_rejected', '?')} quarantined.</p></div>")


def generate(trade_date: str) -> dict:
    rep = build(trade_date)
    stamp = trade_date.replace("-", "")
    out = {"report": rep,
           "excel": excel(rep, REPORT_DIR / f"Market_Report_{stamp}.xlsx"),
           "chart": chart(rep, REPORT_DIR / f"market_chart_{stamp}.png"),
           "html": email_html(rep)}
    (REPORT_DIR / f"market_email_{stamp}.html").write_text(out["html"], encoding="utf-8")
    return out
